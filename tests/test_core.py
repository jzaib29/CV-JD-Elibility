from io import BytesIO
import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter
from pydantic import ValidationError
from applywise import demo
from applywise.documents import InputError, from_file, from_text
from applywise.export import assemble, as_text, as_docx
from applywise.models import Analysis, EditPlan, LayoutSection, Evidence
from applywise.validation import OutputError, validate_analysis, validate_edits, score_analysis


def test_demo_scores():
    cv, jd = demo.documents()
    first = validate_analysis(demo.analysis(), cv, jd, [])
    second = validate_analysis(demo.clarified(), cv, jd, demo.demo_answers(), first)
    assert score_analysis(first).score == 5.8
    assert score_analysis(second).score == 8.1
    assert score_analysis(second).recommendation == "Apply"
    assert not second.questions


def test_gate_unknown_vs_mismatch():
    result = demo.analysis()
    result.requirements[0].status = "not_evidenced"
    assert score_analysis(result).eligibility == "Needs clarification"
    result.requirements[0].status = "contradicted"
    assert score_analysis(result).eligibility == "Confirmed mismatch"


def test_question_limit():
    data = demo.analysis().model_dump()
    data["questions"] *= 4
    with pytest.raises(ValidationError):
        Analysis.model_validate(data)


def test_no_questions_is_valid():
    result = demo.analysis()
    result.questions = []
    validate_analysis(result, *demo.documents(), [])


def test_false_source_rejected():
    result = demo.analysis()
    result.requirements[0].evidence[0].quote = "PhD in Astrophysics"
    with pytest.raises(OutputError):
        validate_analysis(result, *demo.documents(), [])


def test_false_job_requirement_rejected():
    result = demo.analysis()
    result.requirements[0].job_quote = "Ten years of Kubernetes required"
    with pytest.raises(OutputError):
        validate_analysis(result, *demo.documents(), [])


def test_stable_denominator_and_weights():
    cv, jd = demo.documents()
    old, new = demo.analysis(), demo.clarified()
    new.requirements[0].importance = "preferred"
    assert validate_analysis(new, cv, jd, demo.demo_answers(), old).requirements[0].importance == "required"
    new.requirements.pop()
    with pytest.raises(OutputError):
        validate_analysis(new, cv, jd, demo.demo_answers(), old)


def test_fabricated_metrics():
    cv, _ = demo.documents()
    plan = demo.edits()
    plan.edits[0].replacement += " Improved efficiency by 80%."
    with pytest.raises(OutputError):
        validate_edits(plan, cv, [], demo.analysis())


def test_duplicate_edit():
    cv, _ = demo.documents()
    plan = demo.edits()
    plan.edits.append(plan.edits[0])
    with pytest.raises(OutputError):
        validate_edits(plan, cv, [], demo.analysis())


def test_full_layout_cannot_drop_content():
    cv, _ = demo.documents()
    analysis = demo.analysis()
    analysis.restructure = True
    plan = EditPlan(edits=[], layout=[LayoutSection(heading="Resume", block_ids=[b.id for b in cv.blocks])], notes=[])
    validate_edits(plan, cv, [], analysis)
    assert len(assemble(cv, plan, set(), True)) == len(cv.blocks)
    plan.layout[0].block_ids.pop()
    with pytest.raises(OutputError):
        validate_edits(plan, cv, [], analysis)


def test_only_accepted_changes_export():
    cv, _ = demo.documents()
    plan = validate_edits(demo.edits(), cv, [], demo.analysis())
    rows = assemble(cv, plan, {"C5"})
    text = as_text(rows)
    assert plan.edits[1].replacement in text
    assert plan.edits[0].replacement not in text
    assert cv.blocks[2].text in text
    word = WordDocument(BytesIO(as_docx(rows)))
    assert plan.edits[1].replacement in "\n".join(p.text for p in word.paragraphs)


def test_docx_body_order():
    word = WordDocument()
    word.add_paragraph("A computer science graduate with experience building useful student projects in Python.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Python", "pandas"
    word.add_paragraph("Additional information follows the table.")
    stream = BytesIO()
    word.save(stream)
    doc = from_file(stream.getvalue(), "resume.docx")
    assert doc.text.index("Python | pandas") < doc.text.index("Additional")
    assert any("Tables" in warning for warning in doc.warnings)


@pytest.mark.parametrize("text", ["", "tiny", "x" * 24001])
def test_bad_text(text):
    with pytest.raises(InputError):
        from_text(text)


def test_invalid_or_scanned_pdf():
    with pytest.raises(InputError):
        from_file(b"invalid", "resume.pdf")
    writer, stream = PdfWriter(), BytesIO()
    writer.add_blank_page(width=612, height=792)
    writer.write(stream)
    with pytest.raises(InputError, match="little readable text"):
        from_file(stream.getvalue(), "scan.pdf")


def test_answer_evidence():
    cv, _ = demo.documents()
    plan = demo.edits()
    plan.edits[0].evidence.append(Evidence(source_id="A1", quote=demo.demo_answers()[0].text))
    validate_edits(plan, cv, demo.demo_answers(), demo.analysis())
