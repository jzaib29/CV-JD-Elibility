"""Bounded in-memory extraction; no file writes or cross-user cache."""
import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from applywise.models import Block, Document

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 24_000
HEADINGS = {
    "summary", "professional summary", "profile", "objective", "career objective",
    "education", "academic qualifications", "experience", "work experience",
    "professional experience", "internships", "internship", "projects",
    "academic projects", "personal projects", "skills", "technical skills",
    "certifications", "certificates", "achievements", "awards", "publications",
    "volunteering", "volunteer experience", "languages", "interests", "contact",
    "leadership", "extracurricular activities", "references",
}


class InputError(ValueError):
    pass


def from_text(text, name="Pasted text", prefix="C", kind="text", warnings=None):
    text = text.replace("\x00", "").replace("\r\n", "\n").strip()
    if len(text) < 80 or len(text) > MAX_CHARS:
        raise InputError("Provide between 80 and 24,000 characters of readable text per document.")
    blocks, section = [], "Header"
    for line in text.splitlines():
        line = re.sub(r"[ \t]+", " ", line).strip()
        if not line:
            continue
        if prefix == "C" and line.strip("#: ").lower() in HEADINGS:
            section = line.strip("#: ").title()
            continue
        blocks.append(Block(id=f"{prefix}{len(blocks)+1}", section=section, text=line))
    if not 1 <= len(blocks) <= 180:
        raise InputError("Use a document with 1–180 non-empty content lines.")
    warnings = list(warnings or [])
    checks = ["Readable text extracted", "Document size within limits"]
    if kind in {"text", "txt"}:
        warnings.append("Original file layout cannot be assessed from plain text.")
    if prefix == "C":
        if len({b.section for b in blocks}) < 3:
            warnings.append("Few standard headings detected. Review the extracted structure.")
        else:
            checks.append("Standard section headings detected")
        if not re.search(r"[^\s@]+@[^\s@]+\.[^\s@]+", text):
            warnings.append("No email detected. Include one in your final application if appropriate.")
    return Document(name=Path(name).name, kind=kind, blocks=blocks, checks=checks, warnings=warnings)


def from_file(data, filename, prefix="C"):
    if not data or len(data) > MAX_BYTES:
        raise InputError("Use a non-empty file smaller than 5 MB.")
    suffix, warnings = Path(filename).suffix.lower(), []
    try:
        if suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise InputError("This is not a valid PDF.")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise InputError("Remove the PDF password before uploading.")
            if len(reader.pages) > 10:
                raise InputError("Use a PDF of 10 pages or fewer.")
            pages = [p.extract_text() or "" for p in reader.pages]
            if any(len(p.strip()) < 30 for p in pages):
                raise InputError("A PDF page contains little readable text. OCR it first, use DOCX, or paste the complete text.")
            text = "\n".join(pages)
            warnings.append("PDF reading order may differ from its visual layout. Review Source preview.")
        elif suffix == ".docx":
            with ZipFile(BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 25 * 1024 * 1024:
                    raise InputError("The expanded DOCX is too large. Use a simpler document.")
                if "word/document.xml" not in archive.namelist():
                    raise InputError("This is not a valid DOCX.")
            doc = WordDocument(BytesIO(data))
            lines = []
            for item in doc.iter_inner_content():
                if isinstance(item, Paragraph):
                    lines.append(item.text)
                elif isinstance(item, Table):
                    lines.extend(" | ".join(c.text for c in row.cells) for row in item.rows)
            text = "\n".join(lines)
            if doc.tables:
                warnings.append("Tables detected. Check reading order; exports use a single column.")
            if doc.inline_shapes:
                warnings.append("Images are not read. Copy any essential text into the document body.")
            warnings.append("Headers, footers, and floating text boxes are not imported. Put essential information in the body.")
        elif suffix == ".txt":
            text = data.decode("utf-8-sig")
        else:
            raise InputError("Supported files: PDF, DOCX, UTF-8 TXT. Convert legacy DOC to DOCX.")
    except InputError:
        raise
    except Exception as exc:
        raise InputError("The file could not be read. Export a fresh DOCX or paste its text.") from exc
    return from_text(text, filename, prefix, suffix.lstrip("."), warnings)
