from io import BytesIO
from docx import Document as WordDocument
from docx.shared import Inches, Pt, RGBColor


def assemble(cv, plan, accepted, use_layout=False):
    replacements = {e.block_id: e.replacement for e in plan.edits if e.block_id in accepted}
    blocks = {b.id: b for b in cv.blocks}
    if use_layout and plan.layout:
        return [(s.heading, replacements.get(bid, blocks[bid].text))
                for s in plan.layout for bid in s.block_ids]
    return [(b.section, replacements.get(b.id, b.text)) for b in cv.blocks]


def as_text(rows):
    parts, previous = [], None
    for heading, text in rows:
        if heading != previous:
            if parts:
                parts.append("")
            if heading.lower() not in {"header", "contact"}:
                parts.append(heading.upper())
        parts.append(text)
        previous = heading
    return "\n".join(parts)


def as_docx(rows):
    doc = WordDocument()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(.65)
        s.left_margin = s.right_margin = Inches(.75)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("174D43")
    previous = None
    for heading, text in rows:
        if heading != previous and heading.lower() not in {"header", "contact"}:
            doc.add_heading(heading, level=1)
        doc.add_paragraph(text)
        previous = heading
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
